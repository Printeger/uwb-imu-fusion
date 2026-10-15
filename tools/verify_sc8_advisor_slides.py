#!/usr/bin/env python3
"""Audit the actual editable PPTX, exported PDF, frozen numbers and visual receipt."""
import hashlib
import json
import math
import re
import unicodedata
import zipfile
from pathlib import Path
from statistics import mean
from lxml import etree
from pptx import Presentation
import fitz
import yaml

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs/sc8/presentation';CACHE=ROOT/'.cache/sc8-slides'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def norm(s):return re.sub(r'\s+','',unicodedata.normalize('NFKC',s)).replace('−','-').replace('–','-').lower()
def pointer(doc,path):
    v=doc
    for k in path.split('/')[1:]:
        k=k.replace('~1','/').replace('~0','~');v=v[int(k)] if isinstance(v,list) else v[k]
    return v

def verify():
    checks=[]
    def check(name,passed,detail=None):checks.append(dict(name=name,passed=bool(passed),detail=detail))
    req=ROOT/'docs/requirements/sc8_advisor_slides/SC8_ADVISOR_SLIDES_REQUIREMENTS.md'
    story=json.loads((OUT/'STORYBOARD.json').read_text());m=json.loads((ROOT/'docs/sc8/SC8_METRICS.json').read_text());d=json.loads((ROOT/'docs/sc8/SC8_PLOTDATA.json').read_text())
    registry=yaml.safe_load((ROOT/'config/sc8_experiments.yaml').read_text());audit=json.loads((ROOT/'docs/sc8/SC8_AUDIT.json').read_text());figs=json.loads((OUT/'FIGURE_SOURCES.json').read_text())
    pptx=OUT/'SC8_ADVISOR_PROGRESS.pptx';pdf=OUT/'SC8_ADVISOR_PROGRESS.pdf';prs=Presentation(pptx);doc=fitz.open(pdf)
    titles=re.findall(r'^### Slide \d+ — \*\*(.*?)\*\*',req.read_text(),re.M)
    check('requirements_byte_identity',story['requirements_sha256']==sha(req))
    check('storyboard_locked_before_authoring',story.get('locked_before_authoring') is True)
    check('exact_15_main_PPTX_and_PDF_slides',len(prs.slides)==len(doc)==15)
    check('16_by_9_geometry',abs(prs.slide_width/914400-13.333333)<1e-5 and abs(prs.slide_height/914400-7.5)<1e-5)
    check('frozen_result_status',m['status']=='SC8_EVALUATED_WITH_EXECUTION_GAPS')
    check('frozen_computation_identity',registry['frozen']['identity']['head']=='61907ab417450d317e443e4ff10f2cce062deabe')
    check('frozen_metrics_plotdata_and_registry_unchanged',sha(ROOT/'docs/sc8/SC8_METRICS.json')==audit['metrics_sha256'] and sha(ROOT/'docs/sc8/SC8_PLOTDATA.json')==audit['plotdata_sha256'] and sha(ROOT/'config/sc8_experiments.yaml')==audit['registry_sha256'])
    check('original_SC8_figure_assets_unchanged',all(sha(ROOT/'docs/sc8'/p)==h for p,h in audit['artifacts'].items()))
    slide_text=[];minimums=[];matching=[]
    for i,s in enumerate(prs.slides):
        ts=[sh.text for sh in s.shapes if sh.has_text_frame];text='\n'.join(ts);slide_text.append(text)
        check(f'{i+1}: exact storyboard title',norm(titles[i]) in norm(text))
        check(f'{i+1}: exact numbering',f'{i+1}/15' in text)
        check(f'{i+1}: white canvas',str(s.background.fill.fore_color.rgb)=='FFFFFF')
        check(f'{i+1}: takeaway caption and frozen source',all(a in text for a in ['Takeaway:','Source: SC8_RESULTS','61907ab']) and norm(story['slides'][i]['caption']) in norm(text))
        check(f'{i+1}: English visible text and notes',not re.search(r'[\u3400-\u9fff]',text+s.notes_slide.notes_text_frame.text))
        check(f'{i+1}: 2_to_4_English_speaker_bullets',2<=s.notes_slide.notes_text_frame.text.count('• ')<=4)
        check(f'{i+1}: PDF slide geometry',abs(doc[i].rect.width/72-13.333333)<.01 and abs(doc[i].rect.height/72-7.5)<.01)
        pdftext=norm(doc[i].get_text())
        missing=[t for t in ts if norm(t) not in pdftext]
        matching.append(dict(slide=i+1,missing_native_text=missing))
        check(f'{i+1}: PDF contains every native PPTX text box',not missing,missing)
        sizes=[r.font.size.pt for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
        minimums+=sizes;check(f'{i+1}: native_fonts_at_least_9pt',bool(sizes) and min(sizes)>=9)
        outside=[sh.name for sh in s.shapes if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+20 or sh.top+sh.height>prs.slide_height+20]
        check(f'{i+1}: all PPTX shapes within page',not outside,outside)
        check(f'{i+1}: actual PDF preview exists',(CACHE/'preview'/f'slide-{i+1:02d}.png').is_file())
    alltext='\n'.join(slide_text)
    check('E00_through_E08_have_visible_evidence',all(f'E{i:02d}' in alltext for i in range(9)))
    check('slide1_is_editable_architecture_not_an_image',not any(sh.shape_type==13 for sh in prs.slides[0].shapes) and len(prs.slides[0].shapes)>35)
    check('slide1_required_architecture_terms',all(v in slide_text[0].lower() for v in ['rolling certificate','single native commit','fgo','uwb/imu fde','tag pl','conditional','risk horizon']))
    check('one_graph_and_surviving_risk_explained','No second FGO or all-history re-solve.' in slide_text[0] and 'source/risk survive' in slide_text[0])
    check('final_four_specific_decisions',all(norm(v) in norm(slide_text[14]) for v in ['Publication scope','Independent calibration resources','State recovery or execution continuity','Application integrity requirements','hardware validation first','hardware/protocol owner','engineering qualification priority','HAL / VAL, TTA']))
    check('no_generic_thank_you',not re.search(r'thank\s+you',slide_text[-1],re.I))
    check('real_engineering_labels_and_reference_gaps_visible',all(v in alltext for v in ['O-FDE-ENG','O-ENG-NOFDE','Walk physical APE / ARE / RPE = NA','STAR input validation OPEN','GT-assisted','raw unqualified']))
    check('late_FULL_suffix_explicit',all(v in slide_text[12] for v in ['7171/10000','2,829 refused','unverified']))
    check('negative_gains_retained',all(v in slide_text[5] for v in ['-0.007908','-0.067733','-1.13562','-1.23141']))
    check('conditional_receipt_count_not_risk_trial','77,131' in alltext and 'correlated samples ≠ risk trials' in alltext and 'do not establish empirical 10^-6 risk' in alltext)
    numerical=json.loads((OUT/'NUMERICAL_CLAIMS.json').read_text());bad=[]
    for c in numerical['claims']:
        vv=[pointer(m,p) for p in c['paths']];op=c['operation']
        if op=='identity':value=vv[0]
        elif op=='sum':value=sum(vv)
        elif op=='mean':value=mean(vv)
        elif op=='min':value=min(vv)
        elif op=='max':value=max(vv)
        elif op=='percent':value=100*vv[0]
        elif op=='ratio':value=vv[0]/vv[1]
        elif op=='fraction_percent':value=100*vv[0]/vv[1]
        elif op=='kilometers':value=vv[0]/1000
        else:raise AssertionError(op)
        if not math.isclose(value,c['value'],rel_tol=1e-12,abs_tol=1e-12) or format(value,c['format'])+c['suffix']!=c['display']:bad.append(c['label'])
    check('all_numerical_claims_recomputed_from_frozen_JSON',not bad,dict(count=len(numerical['claims']),bad=bad))
    check('claim_source_digest',numerical['metrics_sha256']==sha(ROOT/'docs/sc8/SC8_METRICS.json'))
    check('143_outcomes_and_160_budget',m['registered']==143 and m['logical_run_status_counts']=={'SUCCESS':130,'RESOURCE_LIMIT':12,'FAIL':1} and m['attempt_processes']==159)
    check('one_healthy_point_per_registered_seed',sorted(r['seed'] for r in m['runs'] if r['group']=='E01' and r['stream']=='SC8-H' and r['requested_method']=='O-FULL')==[20261020,20261021,20261022])
    for item in figs['figures']:
        name=item['asset']
        check(name+': frozen input digests',all(sha(ROOT/p)==h for p,h in item['data_digest'].items()))
        check(name+': original SVG provenance',sha(ROOT/'docs/sc8/figures'/item['original_figure'])==item['original_figure_sha256'])
        for ext,v in item['outputs'].items():check(name+': '+ext+' identity',sha(ROOT/v['path'])==v['sha256'])
        svg=etree.parse(str(ROOT/item['outputs']['svg']['path']));fonts=svg.xpath('//*[local-name()="text"]')
        check(name+': editable_SVG_text',bool(fonts))
        font_sizes=[float(n) for t in fonts for n in re.findall(r'font:\s*([\d.]+)px',t.get('style',''))]
        check(name+': projected_chart_fonts_at_least_9pt',bool(font_sizes) and min(font_sizes)>=9,dict(minimum_pt=min(font_sizes) if font_sizes else None))
        check(name+': source_plot_axes_and_legends_not_cropped',item.get('axes_and_legend_bounds_checked') is True)
    native=figs.get('native_visual_sources',[])
    check('native_diagrams_and_tables_have_provenance',sorted(x['slide'] for x in native)==[1,2,14,15])
    for item in native:check('native slide '+str(item['slide'])+': actual source digests',all(sha(ROOT/p)==h for p,h in item['data_digest'].items()))
    for method,axes in d['stanford'].items():
        for axis,views in axes.items():
            for mode,h in views.items():
                check(f'Stanford {method} {axis} {mode}: actual bins',sum(sum(sum(row) for row in a) for a in h['layers'].values())+h['outside']==h['total'])
    with zipfile.ZipFile(pptx) as z:
        ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
        names=set(z.namelist());svgparts=[p for p in names if p.startswith('ppt/media/') and p.endswith('.svg')]
        check('all_12_vector_assets_embedded_with_raster_fallback',len(svgparts)==12)
        missing=[]
        for i in range(1,16):
            path=f'ppt/slides/_rels/slide{i}.xml.rels';tree=etree.fromstring(z.read(path))
            from posixpath import normpath,join
            for rel in tree:
                if rel.get('TargetMode')=='External':continue
                target=normpath(join('ppt/slides',rel.get('Target')))
                if target not in names:missing.append(target)
        check('no_missing_PPTX_relationship_targets',not missing,missing)
    layout=json.loads((OUT/'LAYOUT.json').read_text())
    check('measured_native_text_fits_declared_boxes',all(t['required_height']<=t['h']+.025 for s in layout['slides'] for t in s['text_boxes']))
    receipt=json.loads((CACHE/'render-receipt.json').read_text()) if (CACHE/'render-receipt.json').exists() else {}
    check('PDF_rendered_from_current_actual_PPTX',receipt.get('pptx_sha256')==sha(pptx) and receipt.get('pdf_sha256')==sha(pdf) and receipt.get('pages')==15)
    visual=json.loads((OUT/'VISUAL_REVIEW.json').read_text()) if (OUT/'VISUAL_REVIEW.json').exists() else {}
    check('every_final_slide_visually_reviewed',visual.get('status')=='PASS' and visual.get('pptx_sha256')==sha(pptx) and visual.get('pdf_sha256')==sha(pdf) and sorted(x['slide'] for x in visual.get('slides',[]))==list(range(1,16)))
    check('reviewed_contact_sheet_present',(OUT/'CONTACT_SHEET.png').is_file())
    sources=[ROOT/'tools/sc8_advisor_slides.py',ROOT/'tools/sc8_slides_figures.py',ROOT/'tools/verify_sc8_advisor_slides.py',OUT/'STORYBOARD.json',req]
    result=dict(status='PASS' if all(c['passed'] for c in checks) else 'FAIL',scope='Presentation evidence, editability and PDF parity audit; not a new solver or physical qualification audit.',
        checks=checks,pptx_sha256=sha(pptx),pdf_sha256=sha(pdf),minimum_native_font_pt=min(minimums),numerical_claim_count=len(numerical['claims']),
        sources={str(p.relative_to(ROOT)):sha(p) for p in sources},frozen_result_status=m['status'],render=receipt,visual_review=visual,
        preserved_user_files=['docs/sc8/SC8_RESULTS.html','docs/sc8/SC8_RESULTS.pdf'],pushed=False)
    (OUT/'QA_REPORT.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print('PRESENTATION QA',result['status'],'checks',len(checks),'numeric claims',len(numerical['claims']))
    for c in checks:
        if not c['passed']:print('FAILED',c['name'],c['detail'])
    return result

if __name__=='__main__':
    r=verify();raise SystemExit(0 if r['status']=='PASS' else 1)
