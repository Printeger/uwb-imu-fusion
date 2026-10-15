#!/usr/bin/env python3
"""Build a 15-slide editable SC8 advisor deck from the frozen evidence only.

python tools/sc8_advisor_slides.py build
python tools/sc8_advisor_slides.py render [--soffice /path/to/soffice]
python tools/verify_sc8_advisor_slides.py
"""
import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
from pathlib import Path
from statistics import mean
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
import sc8_slides_figures as f

ROOT=f.ROOT; OUT=f.OUT; CACHE=ROOT/'.cache/sc8-slides'
PPTX=OUT/'SC8_ADVISOR_PROGRESS.pptx'; PDF=OUT/'SC8_ADVISOR_PROGRESS.pdf'
FONT='Liberation Sans';BLUE='216A9B';OFF='8BB7CC';ORANGE='C36A2D';RED='B24F54';GREY='858B91';TEAL='148477';INK='202C34';LINE='CED7DC'
W=13.333333; H=7.5
LAYOUT=[]; CLAIMS=[]

def rgb(h):return RGBColor.from_string(h.lstrip('#'))
def sha(p):return f.sha(p)
def jsonsave(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')

def measured_wrap(s,width,size,bold=False):
    fp='/usr/share/fonts/truetype/liberation/LiberationSans-'+('Bold' if bold else 'Regular')+'.ttf'
    if not Path(fp).exists():fp=subprocess.check_output(['fc-match','-f','%{file}',FONT+(':style=Bold' if bold else '')],text=True)
    font=ImageFont.truetype(fp,round(size*4));limit=width*72*4;lines=[]
    for paragraph in s.split('\n'):
        line=''
        for word in paragraph.split():
            nxt=(line+' '+word).strip()
            if font.getlength(nxt)>limit and line:lines.append(line);line=word
            else:line=nxt
        lines.append(line)
    return '\n'.join(lines)

def tx(slide,s,x,y,w,h,size=15,bold=False,color=INK,align=None,name='text',wrap=True):
    s=str(s);wrapped=measured_wrap(s,w,size,bold) if wrap else s
    lines=wrapped.count('\n')+1
    required=lines*size*1.12/72
    if required>h+.025:raise ValueError(f'Slide {len(LAYOUT)} {name}: needs {required:.3f}in, has {h:.3f}: {wrapped}')
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name=name
    tf=sh.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
    for i,line in enumerate(wrapped.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=Pt(size*1.12)
        if align is not None:p.alignment=align
        for r in p.runs:r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
    LAYOUT[-1]['text_boxes'].append(dict(name=name,text=s,rendered_text=wrapped,x=x,y=y,w=w,h=h,size=size,lines=lines,required_height=required))
    return sh

def box(slide,x,y,w,h,color='FFFFFF',line=LINE,kind=MSO_SHAPE.RECTANGLE):
    sh=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.color.rgb=rgb(line);sh.line.width=Pt(.8)
    effect=sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef')
    if effect is not None:effect.set('idx','0')
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh

def connector(slide,x1,y1,x2,y2,color=GREY,arrow=True):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(1)
    effect=sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef')
    if effect is not None:effect.set('idx','0')
    if arrow:
        head=OxmlElement('a:tailEnd');head.set('type','triangle');head.set('w','sm');head.set('len','sm');sh._element.spPr.get_or_add_ln().append(head)
    return sh

def node(slide,s,x,y,w,h,size=12.5,color=BLUE):
    sh=box(slide,x,y,w,h,line=color);sh.name='editable architecture node: '+s.split('\n')[0]
    # Text is native PPT text, never a diagram screenshot.
    tx(slide,s,x+.08,y+.06,w-.16,h-.10,size=size,bold=True,color=color,name='diagram '+s.split('\n')[0])
    return sh

def table(slide,headers,rows,x,y,widths,rowh=.45,size=13,headh=None):
    headh=headh or rowh;w=sum(widths)
    for j,head in enumerate(headers):tx(slide,head,x+sum(widths[:j])+.04,y+.035,widths[j]-.08,headh-.04,size=size,bold=True,color=BLUE,name='table header')
    connector(slide,x,y+headh,x+w,y+headh,color=LINE,arrow=False)
    yy=y+headh
    for i,row in enumerate(rows):
        for j,value in enumerate(row):tx(slide,str(value),x+sum(widths[:j])+.04,yy+.035,widths[j]-.08,rowh-.045,size=size,name=f'table row {i+1}')
        yy+=rowh;connector(slide,x,yy,x+w,yy,color=LINE,arrow=False)

def picture(slide,name,x,y,w,h):
    png=OUT/'figures'/(name+'.png');svg=OUT/'figures'/(name+'.svg')
    sh=slide.shapes.add_picture(str(png),Inches(x),Inches(y),width=Inches(w),height=Inches(h));sh.name='SC8 evidence '+name
    part=Part(PackURI('/ppt/media/'+name+'.svg'),'image/svg+xml',slide.part.package,svg.read_bytes())
    rid=slide.part.relate_to(part,RT.IMAGE)
    extlist=OxmlElement('a:extLst');ext=OxmlElement('a:ext');ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    from lxml import etree
    sv=etree.Element('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',nsmap={'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
    sv.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',rid);ext.append(sv);extlist.append(ext);sh._element.blipFill.blip.append(extlist)
    LAYOUT[-1]['figures'].append(dict(asset=name,x=x,y=y,w=w,h=h,caption=LAYOUT[-1]['caption']))

def pointer(doc,path):
    a=doc
    for k in path.split('/')[1:]:
        k=k.replace('~1','/').replace('~0','~');a=a[int(k)] if isinstance(a,list) else a[k]
    return a

def calculate(vals,op):
    if op=='identity':return vals[0]
    if op=='sum':return sum(vals)
    if op=='mean':return mean(vals)
    if op=='min':return min(vals)
    if op=='max':return max(vals)
    if op=='percent':return vals[0]*100
    if op=='ratio':return vals[0]/vals[1]
    if op=='fraction_percent':return vals[0]/vals[1]*100
    if op=='kilometers':return vals[0]/1000
    raise ValueError(op)

def build():
    LAYOUT.clear();CLAIMS.clear();m,d=f.load();story=json.loads((OUT/'STORYBOARD.json').read_text());assert len(story['slides'])==15
    figs=json.loads((OUT/'FIGURE_SOURCES.json').read_text());byid={r['id']:r for r in m['runs']};index={r['id']:i for i,r in enumerate(m['runs'])}
    def rp(r,tail):return '/runs/'+str(index[r['id']])+'/'+tail
    def claim(slide,label,paths,fmt='.6g',op='identity',suffix=''):
        if isinstance(paths,str):paths=[paths]
        values=[pointer(m,p) for p in paths];v=calculate(values,op);display=format(v,fmt)+suffix
        CLAIMS.append(dict(slide=slide,label=label,source='docs/sc8/SC8_METRICS.json',paths=paths,operation=op,value=v,format=fmt,suffix=suffix,display=display))
        return display
    def rclaim(slide,label,r,tail,fmt='.6g',**kw):return claim(slide,label,rp(r,tail),fmt,**kw)
    def get(group,stream,method='O-FULL',seed=20261020,window=None):return f.run(m,group,stream,method,seed,window)
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
    prs.core_properties.title='SC8 Advisor Progress: Conditional Rolling Position Integrity'
    prs.core_properties.subject='15-slide evidence-based advisor discussion; frozen solver 61907ab'
    prs.core_properties.author='UWB–IMU Fusion Research';prs.core_properties.keywords='SC8, UWB, IMU, FGO, FDE, conditional integrity'
    healthy=[get('E01','SC8-H',seed=s) for s in [20261020,20261021,20261022]]
    pth=[rp(r,'views/O-FULL/ape_m/rmse') for r in healthy]
    passcount=claim(3,'Healthy timing and timely-spatial profile passes',[rp(r,'tested_profile_deadline_pass') for r in healthy],'d','sum')
    amin=claim(3,'Healthy APE minimum',pth,'.7f','min');amax=claim(3,'Healthy APE maximum',pth,'.7f','max')
    native=[r for r in m['runs'] if r['group']!='E00' and 'axes' in r.get('integrity',{})]
    total=claim(8,'Overall qualified correlated receipts',[rp(r,'integrity/spatial_valid') for r in native],',d','sum')
    ub=claim(8,'Overall H+V observed underbounds',[rp(r,'integrity/axes/'+a+'/underbound') for r in native for a in ['H','V']],'d','sum')
    assert total=='77,131' and ub=='0'
    def new(i,headline):
        spec=story['slides'][i-1];s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('FFFFFF')
        LAYOUT.append(dict(number=i,caption=spec['caption'],text_boxes=[],figures=[]))
        title=spec['title'].replace(': ',':\n',1) if i in [8,13] else spec['title']
        tx(s,title,.5,.30,12.33,.86,size=26,bold=True,name='title')
        tx(s,spec['question'],.5,1.20,12.33,.32,size=16.5,color='52636E',name='research question')
        tx(s,headline,.5,1.66,12.33,.29,size=16.5,bold=True,color=BLUE,name='evidence headline')
        connector(s,.5,2.00,12.83,2.00,color=LINE,arrow=False)
        tx(s,spec['caption'],.5,5.85,12.33,.56,size=11.5,color='44535C',name='evidence caption')
        connector(s,.5,6.48,12.83,6.48,color=LINE,arrow=False)
        tx(s,'Takeaway: '+spec['takeaway'],.5,6.59,12.33,.49,size=15.5,bold=True,name='takeaway')
        scope='T-SIM: conditional tag-position PL' if i in [3,6,7,8,9,10,12,13] else 'T-REAL: empirical, raw unqualified' if i in [4,5] else 'conditional SIM / empirical REAL separated'
        tx(s,f'Source: SC8_RESULTS, {spec["evidence"]}; frozen solver 61907ab; {scope}',.5,7.19,11.5,.18,size=9.5,color=GREY,name='source footer')
        tx(s,f'{i}/15',12.11,7.17,.72,.21,size=10.5,color=GREY,align=PP_ALIGN.RIGHT,name='slide number')
        notes='\n'.join('• '+v for v in spec['speaker_notes'])+'\n\nFrozen solver: 61907ab417450d317e443e4ff10f2cce062deabe. Entry HEAD: fa5b3dc9ffe3c88492352491d79ba022f373bfb5. Result status: SC8_EVALUATED_WITH_EXECUTION_GAPS.\n'+f'Figure provenance: FIGURE_SOURCES.json; metric details: NUMERICAL_CLAIMS.json; slide {i}.'
        s.notes_slide.notes_text_frame.text=notes
        return s
    # 01: editable native nodes, connectors, gate and rolling timeline.
    s=new(1,'20 Hz UWB / 200 Hz IMU  •  default LF/LI = 30/30 epochs  •  committed tag-position PL')
    specs=[('Timestamped raw\nUWB + IMU\nsource IDs / noise',.55,2.20,1.62,.84),('Existing discrete FGO\npre-UWB prepare\nIMU preintegration',2.4,2.20,1.75,.84),('Unified physical\nfault hypothesis\npool',4.42,2.20,1.65,.84),('UWB/IMU FDE\nCandidate testing\nexclude / bridge / repair',6.34,2.20,2.05,.84),('Single native commit\niSAM2 actual tag\n+ commit receipt',8.67,2.20,1.88,.84),('Fault-aware tag PL\nRaw set HPL / VPL\naround committed tag',10.83,2.20,1.96,.84)]
    for a in specs:node(s,*a,size=12)
    for x1,x2 in [(2.17,2.4),(4.15,4.42),(6.07,6.34),(8.39,8.67),(10.55,10.83)]:connector(s,x1,2.62,x2,2.62)
    node(s,'UWB slope RAIM screen',2.4,3.23,2.30,.48,size=12)
    node(s,'Raw IMU fault monitor',2.4,3.87,2.30,.48,size=12)
    connector(s,1.36,3.04,1.36,3.48,arrow=False);connector(s,1.36,3.48,2.4,3.48);connector(s,1.36,3.48,1.36,4.11,arrow=False);connector(s,1.36,4.11,2.4,4.11)
    connector(s,4.70,3.47,5.22,3.04);connector(s,4.70,4.11,5.70,3.04)
    node(s,'Rolling certificate / boundary B_k\np, v, R, ba, bg + historical fault carry\nshared source / time-uniform risk',6.18,3.38,4.37,.88,size=12,color=TEAL)
    connector(s,9.61,3.04,9.61,3.38,color=TEAL)
    connector(s,10.55,3.80,10.70,3.80,color=TEAL,arrow=False);connector(s,10.70,3.80,10.70,2.93,color=TEAL,arrow=False);connector(s,10.70,2.93,10.83,2.93,color=TEAL)
    node(s,'Risk / timeliness / AL\nseparate availability gates',10.83,3.38,1.96,.88,size=12,color=TEAL);connector(s,11.81,3.04,11.81,3.38,color=TEAL)
    tx(s,'T-SIM: conditional raw integrity  |  T-REAL: engineering FDE only; raw unqualified',.55,4.48,12.2,.23,size=13,bold=True,color=TEAL,name='scope strip')
    node(s,'Retired raw epochs',.55,4.92,2.05,.42,size=12,color=GREY);node(s,'Certified boundary B_k',2.9,4.92,2.36,.42,size=12,color=TEAL);node(s,'Explicit active FGO / proof: 30 epochs',5.55,4.92,4.42,.42,size=12);node(s,'New epoch k+1',10.27,4.92,2.52,.42,size=12)
    for a,b in [(2.6,2.9),(5.26,5.55),(9.97,10.27)]:connector(s,a,5.13,b,5.13)
    tx(s,'B_k → B_(k+1): state/error, possible faults and source/risk survive; risk horizon ≠ 30 epochs. No second FGO or all-history re-solve.',.55,5.47,12.2,.36,size=11.5,name='rolling timeline explanation')
    # 02: evidence lanes and permanent six-recording registry.
    counts=[claim(2,k,'/logical_run_status_counts/'+k,'d') for k in ['SUCCESS','RESOURCE_LIMIT','FAIL']]
    registered=claim(2,'Unique registered runs','/registered','d');attempt=claim(2,'Retained execution attempts','/attempt_processes','d')
    s=new(2,f'{registered} registered runs: {counts[0]} SUCCESS | {counts[1]} RESOURCE_LIMIT | {counts[2]} FAIL; {attempt} attempts + 1 ROS control')
    for y,label,source,est,target in [(2.2,'T-SIM','Registered physical law','O-FULL / O-NOFDE','Known truth; conditional PL'),(3.33,'T-REAL','Original recordings','O-FDE-ENG / O-ENG-NOFDE','CAL disclosure; raw unqualified'),(4.46,'T-SEMI','Real + fixed overlay','Same engineering paths','Injected label only; no raw PL')]:
        tx(s,label,.55,y,3.7,.25,size=15,bold=True,color=TEAL);tx(s,source+' → '+est,.55,y+.32,3.7,.43,size=12.5,name='evidence lane');tx(s,'→ '+target,.55,y+.81,3.7,.21,size=12,color=GREY)
    inventory=__import__('yaml').safe_load((ROOT/'config/sc8_experiments.yaml').read_text())['inventory'];dur={r['id']:r['imu_duration_s'] for r in inventory}
    labels=['ISAS-Walk2','ISAS-Walk3','Own Vicon, no obstacle','Own Vicon, obstacle','STAR loop-3d-z_s3','STAR loop-2d-fast_s3']
    statuses=['tracker→IMU/tag OPEN; proxy','same reference OPEN','GT-assisted segmented CAL/TEST','GT-assisted; short recording','gyro/rig input convention OPEN','same input concern']
    rows=[[f'R{i}',labels[i-1],f'{dur[f"R{i}"]:.2f} s',statuses[i-1]] for i in range(1,7)]
    table(s,['ID','Recording','Duration','Physical reference / boundary'],rows,4.47,2.17,[.42,2.65,.88,4.41],rowh=.44,size=12.5,headh=.48)
    tx(s,'SFUISE: official rejection ON; ONLINE causal; FINAL same-run revised history.',4.51,5.43,8.22,.25,size=12.5,color=GREY,name='baseline scope')
    # 03: three actual seeds, distinct online / final support.
    s=new(3,f'O-FULL tag APE RMSE {amin}–{amax} m; default timing / timely-spatial profile passes {passcount}/3')
    picture(s,'s03_healthy',.50,2.15,7.55,3.55)
    arpaths=[rp(r,'views/O-FULL/are_rad/rmse') for r in healthy]
    aremin=claim(3,'Healthy ARE min',arpaths,'.4f','min');aremax=claim(3,'Healthy ARE max',arpaths,'.4f','max')
    r=healthy[0];core=rclaim(3,'Healthy seed20 core P95',r,'timing/core_ms/p95','.2f');sink=rclaim(3,'Healthy seed20 sink P95',r,'timing/measurement_to_sink_ms/p95','.2f')
    table(s,['Full path / model','Observed result'],[['Tag APE RMSE (m)',amin+'–'+amax],['Body ARE RMSE',aremin+'–'+aremax+' rad'],['Seed20 core / sink P95',core+' / '+sink+' ms'],['Study limits','40 / 50 ms; timely ≥95%']],8.38,2.32,[2.38,2.07],rowh=.50,size=13,headh=.40)
    tx(s,'Per-run output-stamp scoring.\nFULL / NOFDE: 600 states each.\nSFUISE: 300 online; 298 valid final.\nFinal history is not an online gain.',8.42,4.90,4.30,.78,size=12,name='healthy support')
    # 04: empirical example with physical interpretation boundaries.
    s=new(4,'Walk physical APE / ARE / RPE = NA; STAR input validation OPEN; own Vicon is GT-assisted')
    picture(s,'s04_vicon',.5,2.15,7.50,3.55)
    table(s,['Family','Usable evidence','Open boundary'],[['Walk R1/R2','Time coverage; proxy diagnostic','Tracker–body/tag reference'],['Own R3/R4','Empirical TEST; CAL disclosed','TEST-involved fit; not blind'],['STAR R5/R6','Frozen-interface behavior','Gyro/rig convention; no intrinsic ranking']],8.35,2.23,[1.0,1.66,1.80],rowh=.79,size=13,headh=.55)
    star=get('E02','R5','SFUISE');km=claim(4,'STAR R5 frozen-interface APE km',rp(star,'views/S-ONLINE/ape_m/rmse'),'.2f','kilometers')
    tx(s,'R3 FULL: no complete TEST trajectory.\nSTAR R5 S-ONLINE APE: '+km+' km.\nInput OPEN: no intrinsic-method ranking.',8.39,5.23,4.34,.55,size=12,name='real limitations')
    # 05: temporal coverage, original denominators retained.
    s=new(5,'4/6 O-FDE-ENG real runs hit the resource limit; every real raw PL remains unqualified')
    picture(s,'s05_coverage',.5,2.2,7.6,3.45)
    covrows=[]
    for i in range(1,7):
        row=[f'R{i}']
        for method in ['O-FULL','O-NOFDE']:
            r=get('E02',f'R{i}',method);score=next(iter(r['views'].values()));c=score.get('output_coverage_on_GT_TEST_grid')
            if c is None:row.append('NA');continue
            key=next(iter(r['views']));tail='views/'+key+'/output_coverage_on_GT_TEST_grid/'
            n=rclaim(5,'R'+str(i)+' '+method+' covered',r,tail+'covered','d');den=rclaim(5,'R'+str(i)+' '+method+' planned',r,tail+'planned','d')
            claim(5,'R'+str(i)+' '+method+' coverage percent',rp(r,tail+'fraction'),'.2f',op='percent',suffix='%')
            row.append(n+'/'+den)
        covrows.append(row)
    table(s,['GT grid','FDE covered/N','NoFDE covered/N'],covrows,8.38,2.24,[.7,1.84,1.91],rowh=.41,size=12.5,headh=.45)
    tx(s,'Walk: temporal proxy grid only.\nR3 FULL has no saved TEST output.\nRates differ from commit/attempt rate.',8.42,5.17,4.28,.58,size=12,name='coverage interpretation')
    # 06: actual paired gains and signed seed means.
    s=new(6,'OFF − ON: positive is improvement; negative is degradation; all three seeds remain visible')
    picture(s,'s06_gains',.5,2.20,8.1,3.45);rows=[]
    for stream,label in [('SC8-U','UWB'),('SC8-A','Accel'),('SC8-G','Gyro'),('SC8-J','Joint')]:
        pp=[(i,p) for i,p in enumerate(m['paired']) if p['group']=='E03' and p['stream']==stream]
        av=claim(6,label+' mean APE gain',[f'/paired/{i}/gains/ape_m/off_minus_on' for i,p in pp],'+.6f','mean')
        rv=claim(6,label+' mean ARE gain',[f'/paired/{i}/gains/are_rad/off_minus_on' for i,p in pp],'+.5f','mean');rows.append([label,av,rv])
    table(s,['Fault','ΔAPE (m)','ΔARE (rad)'],rows,8.88,2.32,[.85,1.48,1.52],rowh=.43,size=13,headh=.40)
    tx(s,'Black ticks are seed means.\nAccel: better position, worse attitude.\nGyro: both means deteriorate.\nJoint: attitude gain varies by seed.',8.92,4.78,3.80,.86,size=13,name='fault interpretation')
    # 07: actual function events do not imply accuracy recovery.
    r=get('E03','SC8-J');imu=r['episodes'][0];s=new(7,'T-SIM SC8-J / seed 20261020: actual prior repair + bridge + rejoined receipts; accuracy is separate')
    picture(s,'s07_fault_timeline',.5,2.12,8.90,3.60)
    table(s,['Source / event','Magnitude / receipt'],[['UWB bias','+2.25 m / 3 epochs'],['Raw accel','+20 m/s² / 500 ms'],['First FD / bridge','Epoch 121 / 121'],['First rejoin4','Epochs 141–144']],9.68,2.27,[1.25,1.85],rowh=.55,size=12.5,headh=.48)
    tx(s,'Burgundy: accel pulse.\nTeal: UWB pulse.\nRejoined is a persistent state,\nnot an accuracy certificate.',9.72,5.02,3.01,.75,size=12,name='fault bands')
    for name,tail,fmt in [('Fault accel amplitude','episodes/0/amplitude','.0f'),('Fault first bridge','episodes/0/first_bridge','d'),('Fault UWB amplitude','episodes/1/amplitude','.2f')]:rclaim(7,name,r,tail,fmt)
    # 08: same actual histograms, no synthetic scatter from averages.
    s=new(8,f'{total} qualified T-SIM receipts overall; observed H/V underbounds = {ub}; correlated samples ≠ risk trials')
    picture(s,'s08_horizontal',.55,2.12,12.2,3.55)
    # 09: vertical and diagnostic AL sensitivity, all three healthy seeds.
    s=new(9,'Same estimates and PL: research AL (2,3) m gives 3.5–5.667%; (5,5) m gives 99.67–100%')
    picture(s,'s09_vertical',.50,2.12,8.55,3.55);picture(s,'s09_availability',9.48,2.16,3.25,1.90)
    rows=[]
    for al,fmt in [('2/3','.3f'),('5/5','.2f')]:
        vals=[]
        for r in healthy:vals.append(claim(9,al+' seed'+str(r['seed'])+' availability',[rp(r,'integrity/joint_AL/'+al.replace('/','~1')+'/timely_within_AL'),rp(r,'integrity/planned')],fmt,'fraction_percent',suffix='%'))
        raw=[float(v[:-1]) for v in vals];rows.append([al+' m',f'{min(raw):g}–{max(raw):g}%'])
    table(s,['HAL / VAL','Joint timely'],rows,9.32,4.07,[1.5,2.0],rowh=.32,size=12.5,headh=.32)
    tx(s,'O-FULL; 3 independent seeds.\nDenominator: all 600 epochs / seed.\nNo application certification.',9.36,5.18,3.4,.56,size=12,name='AL denominator')
    # 10: window quality and cost on exact common 1000-epoch interval.
    s=new(10,'7 unique configurations × 2 streams; common epochs 1101–2100; default core study target = 40 ms')
    picture(s,'s10_window',.55,2.12,12.2,2.45)
    wh30=get('E05','SC8-WH',window=[30,30]);wh1k=get('E05','SC8-WH',window=[1000,1000]);wj30=get('E05','SC8-WJ',window=[30,30]);wj1k=get('E05','SC8-WJ',window=[1000,1000])
    rows=[]
    for label,a,b,tail,fmt,interp in [('Healthy APE (m)',wh30,wh1k,'views/O-FULL/ape_m/rmse','.8f','No meaningful gain'),('Healthy core P95 (ms)',wh30,wh1k,'timing/core_ms/p95','.3f','Higher measured cost'),('Fault APE (m)',wj30,wj1k,'views/O-FULL/ape_m/rmse','.8f','Benefit in this stream')]:
        rows.append([label,rclaim(10,label+' 30/30',a,tail,fmt),rclaim(10,label+' 1000/1000',b,tail,fmt),interp])
    table(s,['Same common interval','30/30','1000/1000','Interpretation'],rows,.55,4.62,[3.5,2.2,2.2,4.3],rowh=.27,size=12.5,headh=.29)
    # 11: distinguish registered main healthy and E06 performance-repeat joint.
    s=new(11,'Core P95 ≤40 ms + sink P95 ≤50 ms + timely spatial ≥95%; two E06 serial confirmations pass')
    picture(s,'s11_latency',.55,2.10,12.2,2.5)
    rows=[]
    for label,r in [('E01 H / seed20 FULL',healthy[0]),('E06 J / seed20 FULL',get('E06','SC8-J')),('E01 H / seed20 NOM',get('E01','SC8-H','O-NOM'))]:
        rows.append([label,rclaim(11,label+' core',r,'timing/core_ms/p95','.3f'),rclaim(11,label+' sink',r,'timing/measurement_to_sink_ms/p95','.3f'),'Full service' if r['requested_method']=='O-FULL' else 'Nominal only; no matched PL'])
    real=[r for r in m['runs'] if r['group']=='E02' and r['requested_method']=='O-FULL']
    rrmax=claim(11,'Real engineering max measurable core P95',[rp(r,'timing/core_ms/p95') for r in real],'.1f','max')
    sinkmax=claim(11,'Real engineering max measurable sink P95',[rp(r,'timing/measurement_to_sink_ms/p95') for r in real],'.3g','max')
    rows.append(['E02 real O-FDE-ENG','up to '+rrmax,'up to '+sinkmax,'Measured prefix; 4/6 resource-limited'])
    table(s,['Run / capability','Core P95 (ms)','Sink P95 (ms)','Boundary'],rows,.55,4.60,[3.5,2.2,2.2,4.3],rowh=.24,size=12,headh=.23)
    # 12: legal ablations, separate M/A/P semantics.
    s=new(12,'Removing support/joint constraints widens PL; the old G2 numerical candidate commits 0/80')
    picture(s,'s12_ablation',.5,2.12,8.1,3.55)
    for y,title,body in [(2.23,'M — same center','Old/new PL at one committed tag center: a tightening witness.'),(3.17,'A — full trajectory','Changing fault support can change the subsequent state trajectory.'),(4.13,'P — performance','Whole-service measurements; do not sum module percentiles.')]:
        tx(s,title,8.95,y,3.82,.26,size=15,bold=True,color=TEAL);tx(s,body,8.95,y+.33,3.82,.51,size=14)
    g2=get('E07','SC8-G2','O-ABL-G2');abort=rclaim(12,'G2 measured abort core',g2,'timing/core_ms/p95','.3f')
    tx(s,'G2: 0/80; epoch1 abort = '+abort+' ms.\nAPE / PL: NA; no common valid baseline.\nJoint + support is a combined switch.',8.95,5.05,3.82,.62,size=12.5,bold=True,color=RED)
    for r in [r for r in m['runs'] if r['group']=='E07']:
        for axis in ['H','V']:
            tail='integrity/axes/'+axis+'/pl/mean'
            if pointer(m,rp(r,tail)) is not None:rclaim(12,r['id']+' '+axis+'PL',r,tail,'.5f')
    # 13: complete L0, partial FULL L1 and complete NOFDE L1.
    s=new(13,'All three processes exit SUCCESS; actual FULL L1 coverage is 7171/10000, not full completion')
    picture(s,'s13_long',.5,2.12,8.75,3.55);rows=[]
    for stream,method,label in [('SC8-L0','O-FULL','L0 FULL'),('SC8-L1','O-FULL','L1 FULL'),('SC8-L1','O-NOFDE','L1 NOFDE')]:
        r=get('E08',stream,method);c=rclaim(13,label+' committed',r,'max_committed_epoch','d');n=rclaim(13,label+' planned',r,'planned_epochs','d');q=rclaim(13,label+' queue peak',r,'source_delivery_counts/queue_peak','d');rows.append([label,c+'/'+n,q])
    table(s,['Run','Committed / plan','Queue peak'],rows,9.53,2.35,[1.02,1.40,.88],rowh=.58,size=12.5,headh=.5)
    tx(s,'FULL L1: 2,829 refused epochs.\nThe final scheduled fault is unverified.\nSlots grow; bounded active window\ndoes not mean constant total memory.',9.57,4.95,3.21,.76,size=12.5,name='long-stream caveat')
    # 14: explicit evidence-to-claim matrix, including omissions.
    s=new(14,'Demonstrated: conditional method + measured simulation results; not hardware-qualified real integrity')
    rows=[['Conditional rolling position PL','77,131 qualified receipts; 0 observed underbounds','Registered source/noise/fault law only; correlated observations'],['Healthy tag positioning ≈5 cm','3 fixed independent simulation seeds','Different initialization information from SFUISE'],['FDE benefit is fault-specific','Accel/joint position gains; gyro regression','Attitude / late recovery can deteriorate'],['Real comparison remains incomplete','Walk/STAR interface gaps; real FDE censoring','Engineering empirical results; no raw real qualification'],['Default small window justified','25.096 → 795.501 ms healthy core P95','Longer lag helps this fault stream; not generally real-time']]
    table(s,['Defensible claim','Direct SC8 evidence','Boundary / not established'],rows,.55,2.16,[3.70,4.0,4.50],rowh=.57,size=13.5,headh=.40)
    tx(s,'Additional gaps: SFUISE velocity/bias export; invalid history prefix; weak-UWB and S1/S2 studies omitted.',.59,5.52,12.1,.23,size=12,color=GREY,name='remaining gaps')
    # 15: exactly the four actionable advisor questions; no generic thank-you.
    s=new(15,'Proposed next action: close independent calibration → stabilize real FDE → limited newly frozen rerun')
    cards=[('1. Publication scope','Should the near-term paper focus on the conditional rolling raw position PL + joint FDE architecture, with honest execution/real-data limitations, or is physically qualified real-data integrity required before submission?','Decision: theory + qualified simulation first, or hardware validation first.'),('2. Independent calibration resources','Can we obtain independent tracker↔IMU/body↔tag 6-DoF extrinsics, timing and anchor-world references, or collect a new calibration sequence with three-axis motion and independent held-out TEST?','Decision: identify the hardware/protocol owner and available sensors.'),('3. State recovery or execution continuity','Should we prioritize explaining and fixing gyro-induced attitude degradation / long-term state recovery, or first make engineering FDE process complete real recordings within a bounded resource budget?','Decision: algorithmic research priority or engineering qualification priority.'),('4. Application integrity requirements','For the intended deployment, what HAL / VAL, TTA, covered fault classes and failure action should determine meaningful protected availability?','Decision: target application and first defensible safety requirements.')]
    for i,(head,q,decision) in enumerate(cards):
        x=.55+(i%2)*6.23;y=2.20+(i//2)*1.77
        tx(s,head,x,y,5.97,.27,size=16,bold=True,color=BLUE,name='advisor question heading')
        tx(s,q,x,y+.35,5.97,.96,size=15,name='advisor question')
        tx(s,decision,x,y+1.36,5.97,.40,size=12.5,bold=True,color=TEAL,name='requested decision')
    assert len(prs.slides)==15;prs.save(PPTX)
    jsonsave(OUT/'NUMERICAL_CLAIMS.json',dict(schema='sc8-slide-claims-v1',metrics_sha256=sha(ROOT/'docs/sc8/SC8_METRICS.json'),claims=CLAIMS))
    jsonsave(OUT/'LAYOUT.json',dict(slides=LAYOUT,font=FONT,geometry_inches=[W,H],minimum_native_font=9.5))
    # English source manifest is generated directly from immutable provenance.
    lines=['# SC8 advisor presentation: figure sources','', 'Main deck: exactly 15 slides. Frozen solver: `61907ab417450d317e443e4ff10f2cce062deabe`.', '',
        'Charts are rearranged from the original frozen SC8 SVG/summary evidence, with new SVG assets and 300-dpi PNG fallbacks embedded in PPTX. All native architecture nodes, arrows, tables, takeaways and advisor questions remain editable.', '',
        'The report and plotdata are read-only. `FIGURE_SOURCES.json` records input and output hashes; `NUMERICAL_CLAIMS.json` records exact paths, arithmetic and formatting for numerical claims. A process SUCCESS is never interpreted as full trajectory completion.', '',
        '| Slide | Asset / original evidence | Section and displayed metric | Caveat |','|---|---|---|---|']
    for item in figs['figures']:lines.append('| '+str(item['slide'])+' | '+item['asset']+' / '+item['original_figure']+' | '+item['section']+' | '+item['caveat']+' |')
    lines+=['','Slides 1, 2, 14 and 15 use native editable diagrams / tables. Architecture follows `research/sc6_service.cpp`, `research/rolling_policy.hpp`, and the verified SC7/SC8 descriptions. Historical reports supply architecture only, never SC8 performance values.', '',
        'Reproduce (Python 3.8+; Liberation Sans or installed metric-compatible font):','', '```bash',
        'python -m pip install -r docs/sc8/presentation/requirements.txt',
        'python tools/sc8_slides_figures.py',
        'python tools/sc8_advisor_slides.py build',
        'python tools/sc8_advisor_slides.py render --soffice /path/to/soffice',
        'python tools/verify_sc8_advisor_slides.py', '```','',
        'PDF is exported from the actual PPTX with LibreOffice Impress. The renderer also writes all 15 PNG previews and a contact sheet under ignored `.cache/sc8-slides/preview/`. It launches no experimental solver.','',
        'The referenced GOAL_SC8_ADVISOR_SLIDES.txt was absent from the supplied directory. The complete provided user objective and SC8_ADVISOR_SLIDES_REQUIREMENTS.md define the locked 15-slide storyboard; no replacement goal was fabricated.']
    (OUT/'FIGURE_SOURCES.md').write_text('\n'.join(lines)+'\n')
    print('PPTX built:',PPTX,'numerical claims',len(CLAIMS),flush=True)

def render(soffice=None):
    assert PPTX.is_file(), 'Build PPTX before exporting'
    CACHE.mkdir(parents=True,exist_ok=True)
    if PDF.exists():PDF.unlink()
    local=CACHE/'lo/usr/lib/libreoffice/program/soffice'
    office=Path(soffice or shutil.which('libreoffice') or shutil.which('soffice') or local).resolve()
    if not office.is_file():raise RuntimeError('LibreOffice Impress required for actual PPTX export; pass --soffice')
    env=dict(os.environ);args=[str(office),'--headless','--nologo','--nodefault','--nofirststartwizard']
    if str(office).startswith(str(CACHE/'lo')):
        lib=CACHE/'lo/usr/lib';env['LD_LIBRARY_PATH']=str(lib/'x86_64-linux-gnu')+':'+str(lib/'libreoffice/program')+':'+env.get('LD_LIBRARY_PATH','')
        args+=['-env:BRAND_BASE_DIR='+(lib/'libreoffice').as_uri()]
    env['SAL_USE_VCLPLUGIN']='svp';args+=['-env:UserInstallation='+(CACHE/'lo-profile').as_uri(),'--convert-to','pdf:impress_pdf_Export','--outdir',str(OUT),str(PPTX)]
    r=subprocess.run(args,env=env,capture_output=True,text=True,timeout=180)
    print(r.stdout,r.stderr,flush=True);assert r.returncode==0 and PDF.exists(),'PPTX PDF export failed'
    import fitz
    from PIL import Image,ImageOps,ImageDraw
    doc=fitz.open(PDF);assert len(doc)==15
    prev=CACHE/'preview';prev.mkdir(exist_ok=True);thumbs=[]
    for i,p in enumerate(doc):
        im=Image.open(__import__('io').BytesIO(p.get_pixmap(matrix=fitz.Matrix(1.65,1.65),alpha=False).tobytes('png'))).convert('RGB');im.save(prev/f'slide-{i+1:02d}.png')
        im.thumbnail((640,360));cell=Image.new('RGB',(660,394),'#e9edef');cell.paste(im,((660-im.width)//2,8));ImageDraw.Draw(cell).text((12,373),f'Slide {i+1}/15',fill='#202c34');thumbs.append(cell)
    sheet=Image.new('RGB',(660*3,394*5),'white')
    for i,im in enumerate(thumbs):sheet.paste(im,((i%3)*660,(i//3)*394))
    sheet.save(prev/'contact-sheet.png')
    jsonsave(CACHE/'render-receipt.json',dict(renderer=str(office),pptx_sha256=sha(PPTX),pdf_sha256=sha(PDF),pages=len(doc),mode='Actual PPTX → LibreOffice Impress PDF → PyMuPDF slide previews'))
    print('Exported actual PPTX to PDF; rendered 15 slide previews and contact sheet',flush=True)

if __name__=='__main__':
    p=argparse.ArgumentParser(__doc__);p.add_argument('command',choices=['figures','build','render']);p.add_argument('--soffice');a=p.parse_args()
    if a.command=='figures':f.figures()
    elif a.command=='build':build()
    else:render(a.soffice)
